"""SAC-2 精度基准测试（PRD：与手工验算偏差 < 0.1%）

从 PRD 测试数据 Excel 导入功能区域参数，执行静态计算，
验证末端负荷与手工公式计算结果一致（偏差 < 0.1%）。
"""
import os
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.calculation.services import calc_terminal_load
from apps.projects.models import Room

TEST_DATA_DIR = os.path.join(
    os.path.dirname(__file__), "..", "..", "..", "..",
    "06-测试文档", "测试数据",
)


def load_excel_rooms(filename: str) -> list[dict]:
    """从 Excel 加载功能区域数据（PRD 附录A 列映射）"""
    filepath = os.path.join(TEST_DATA_DIR, filename)
    if not os.path.exists(filepath):
        pytest.skip(f"测试数据文件不存在: {filename}")
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]

    # 构建列映射
    col_map = {}
    for i, h in enumerate(headers):
        col_map[str(h).strip()] = i

    rooms = []
    for row in rows[1:]:
        if not row or not row[col_map.get("房间名称", 1)]:
            continue
        rooms.append({
            "room_name": str(row[col_map["房间名称"]]),
            "area": Decimal(str(row[col_map["面积(m2)"]])),
            "height": Decimal(str(row[col_map["吊顶高度(m)"]])),
            "civil_load_index": Decimal(str(row[col_map["土建指标(W/m2)"]])),
            "lighting_load_index": Decimal(str(row[col_map.get("照明指标(W/m2)", -1)])) if col_map.get("照明指标(W/m2)") and row[col_map["照明指标(W/m2)"]] else Decimal("15"),
            "personnel_load_index": Decimal(str(row[col_map.get("人员指标(W/人)", -1)])) if col_map.get("人员指标(W/人)") and row[col_map.get("人员指标(W/人)")] and row[col_map["人员指标(W/人)"]] else Decimal("0"),
            "personnel_count": int(row[col_map.get("人数", -1)]) if col_map.get("人数") and row[col_map.get("人数")] and row[col_map["人数"]] else 0,
        })
    return rooms


def manual_calc_terminal(room_data: dict) -> Decimal:
    """手工计算末端负荷（PRD §15.1，作为基准值）

    土建 = civil × area / 1000
    照明 = lighting × area / 1000
    人员 = personnel_idx × count / 1000
    末端 = 土建 + 照明 + 人员（简化：不含设备，Excel 可能缺设备列）
    """
    civil = room_data["civil_load_index"] * room_data["area"] / Decimal("1000")
    lighting = room_data["lighting_load_index"] * room_data["area"] / Decimal("1000")
    personnel = room_data["personnel_load_index"] * room_data["personnel_count"] / Decimal("1000")
    return civil + lighting + personnel


@pytest.mark.django_db
class TestPrecisionSAC2:
    """SAC-2: 静态冷量计算偏差 < 0.1%"""

    @pytest.mark.parametrize("filename", [
        "PCB工厂测试数据_西北_西安.xlsx",
        "PCB工厂测试数据_广东_广州.xlsx",
        "PCB工厂测试数据_东南亚_胡志明市.xlsx",
    ])
    def test_terminal_load_precision(self, filename, floor):
        """末端负荷计算与手工验算偏差 < 0.1%（SAC-2）"""
        rooms_data = load_excel_rooms(filename)
        assert len(rooms_data) > 0, f"{filename} 无有效数据"

        for rd in rooms_data:
            # 创建 Room 实例（不 save，仅用于计算）
            room = Room(
                floor=floor,
                room_name=rd["room_name"],
                area=rd["area"],
                height=rd["height"],
                civil_load_index=rd["civil_load_index"],
                lighting_load_index=rd["lighting_load_index"],
                personnel_load_index=rd["personnel_load_index"],
                personnel_count=rd["personnel_count"],
            )
            # 系统计算
            result = calc_terminal_load(room)
            system_value = result["civil"] + result["lighting"] + result["personnel"]

            # 手工基准值
            manual_value = manual_calc_terminal(rd)

            if manual_value > 0:
                deviation = abs(system_value - manual_value) / manual_value * Decimal("100")
                assert deviation < Decimal("0.1"), (
                    f"{rd['room_name']}: 系统={system_value} 手工={manual_value} "
                    f"偏差={deviation}% >= 0.1%"
                )

    def test_civil_load_exact_match(self, floor):
        """土建负荷精确匹配：civil_index × area / 1000"""
        rooms_data = load_excel_rooms("PCB工厂测试数据_西北_西安.xlsx")
        for rd in rooms_data[:5]:  # 前 5 个验证
            room = Room(
                floor=floor, room_name=rd["room_name"],
                area=rd["area"], height=rd["height"],
                civil_load_index=rd["civil_load_index"],
                lighting_load_index=rd["lighting_load_index"],
            )
            result = calc_terminal_load(room)
            expected = rd["civil_load_index"] * rd["area"] / Decimal("1000")
            assert abs(result["civil"] - expected) < Decimal("0.0001"), (
                f"{rd['room_name']} 土建负荷: {result['civil']} vs {expected}"
            )
